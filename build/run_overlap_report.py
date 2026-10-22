# -*- coding: utf-8 -*-
"""러닝 겹침 점검표(엑셀) — 26SS 러닝과 26FW 러닝(시점재고 합산)이 같은 구품번을 두 번 세는 것을 찾아 뺀 근거.
build/run_annual_snapshot.json(빌드 때 run_annual.py 가 만든 값)을 읽어 바탕화면에 '26FW_러닝_겹침_점검.xlsx' 를 쓴다.
사용:  python build/run_overlap_report.py [저장경로]"""
import json
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.expanduser('~'), 'Desktop', '26FW_러닝_겹침_점검.xlsx')
R = json.load(open(os.path.join(HERE, 'run_annual_snapshot.json'), encoding='utf-8'))
ck = R['checks']
TOL = ck['tol']
wb = openpyxl.Workbook()
HEAD = PatternFill('solid', fgColor='1F2937')


def sheet(ws, title, header, rows, widths):
    ws.title = title
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = HEAD
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'B2'
    ws.row_dimensions[1].height = 34
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if isinstance(c.value, (int, float)):
                c.number_format = '#,##0'
    return ws


ov1, ov2, only, allss = R['overlap'], R['overlapNew'], R['fwOnly'], R['ssAll']
c_all, c1, c2 = ck['carryAll'], ck['carryOverlap'], ck['carryOverlapNew']
ws = wb.active
ws.title = '요약'
lines = [
    ['26FW · 26SS 러닝 겹침 점검', ''],
    ['', ''],
    ['무엇을 점검했나', '26SS 러닝(1/1~)과 26FW 러닝 ‘시점재고 합산’(7/31 시점재고 + 8/1~ 판매)이 같은 구품번을 두 번 세는지, 두 스타일맵이 서로 맞는지(오차 %d장 이내)' % TOL],
    ['SS 러닝 품번', ck['ssN']],
    ['FW 러닝 구품번(이월)', ck['fwOldN']],
    ['① SS 러닝과 겹치는 구품번(W 코드)', ck['overlapN']],
    ['② 26SS 신상품과 겹치는 구품번(GF1 코드)', ck['overlapNewN']],
    ['③ SS 러닝 목록 밖 FW 구품번(구품번 추가, SS 기준으로 더함)', ck['onlyFwN']],
    ['④ 과년 러닝(러닝 원본 NAS 24SS~25FW 목록에만 있음, 구품번 추가·SS 기준으로 더함 — MD 결정 10/22)', ck['onlyExtraN']],
    ['  구품번 추가 합계(③+④)', ck['onlyN']],
    ['  ④ 과년 러닝이 올라 있는 시즌 목록', '%s (26SS·26FW 목록에도 있는 품번: %s · 러닝 원본에서 비어 있는 시즌 열: %s)' % (', '.join(ck.get('extraSeasons') or []), ', '.join(ck.get('extraIn26') or []) or '없음', ', '.join(ck.get('nasEmptySeasons') or []) or '없음')],
    ['러닝 원본(NAS 런닝품번리스트) 품번 반영 현황', '%d개 = SS 러닝 %d + FW 구품번 %d + 구품번 추가(④) %d · 미반영 %s' % (
        (ck.get('nasCover') or {}).get('n', 0), (ck.get('nasCover') or {}).get('ss', 0), (ck.get('nasCover') or {}).get('fwOld', 0), (ck.get('nasCover') or {}).get('extra', 0),
        ', '.join((ck.get('nasCover') or {}).get('none') or []) or '없음')],
    ['기준', '시즌 지표로 보면 구품번 = 26SS(1/1~ 지금까지), 신품번(GF7·GF8)만 = 26FW(8/1~). FW ‘시점재고 합산’의 이월(구품번 36개)은 연간에 더하지 않고, 구품번은 SS 기준으로 한 번만 센다.'],
    ['', ''],
    ['FW 이월 합계 — 7/31 시점재고(장)', c_all['qtyIn']],
    ['  ① 겹침분', c1['qtyIn']],
    ['  ② 겹침분', c2['qtyIn']],
    ['  ③ 나머지(연간에서는 SS 기준으로 대신 센다)', c_all['qtyIn'] - c1['qtyIn'] - c2['qtyIn']],
    ['FW 이월 합계 — 8/1~ 판매(장)', c_all['sales']],
    ['  ① 겹침분', c1['sales']],
    ['  ② 겹침분', c2['sales']],
    ['  ③ 나머지', c_all['sales'] - c1['sales'] - c2['sales']],
    ['', ''],
    ['대조 결과(허용 오차 %d장)' % TOL, ''],
    ['주간판매(52주) 최대 차이 — SS 러닝 %d품번' % ck['ssN'], ck['weeklyMaxDiffAll']],
    ['주간판매(52주) 최대 차이 — 26SS 신상품과 겹침 %d품번' % ck['overlapNewN'], ck['weeklyMaxDiffNew']],
    ['현재고 최대 차이(SS 러닝 / 신상품 겹침)', '%s / %s' % (int(ck['stockMaxDiffAll']), int(ck['stockMaxDiffNew']))],
    ['FW 이월 판매(8/1~)가 SS 주간판매 경계 범위(8/3주~ ~ 7/27주~) 밖인 품번', ', '.join(ck['boundaryFail'] + ck['boundaryFailNew']) or '없음'],
    ['FW 이월 계산 = 종합 시트(CJ~CR)', '일치' if ck.get('carryMatchesSheet') else '불일치'],
    ['SS 러닝 품번이 26SS 신상품에도 있는 경우', ', '.join(ck.get('ssRunInNew') or []) or '없음'],
    ['', ''],
    ['이어붙임(SS 러닝과 겹치는 %d개의 1/1~ 판매)' % ck['overlapN'], ''],
    ['스타일맵 정의(런닝 BACK 판매량 = 올해재고 − 현재고)', ck['overlapSsSales']],
    ['판매기록 기준(런닝 주마감 누계 − 12월기초재고 누계)', ck['overlapSsSalesRec']],
    ['차이(품번별 최대 / %d장 넘는 품번)' % TOL, '%s장 / %s' % (int(ck['ssSalesDiffMax']), ', '.join(ck['ssSalesDiffOver']) or '없음')],
    ['SS 러닝 31품번 전체: 스타일맵 정의 / 판매기록', '%s / %s' % (int(ck['ssSalesBackTotal']), int(ck['ssSalesRecTotal']))],
    ['해석', '두 파일은 같은 ERP 추출이라 판매기록은 일치. SS 스타일맵의 러닝 판매량은 올해재고−현재고로 계산해 신규발주와 실제 입고의 차이만큼 판매기록보다 적다. 연간에는 스타일맵 정의 그대로 사용.'],
    ['', ''],
    ['연간 러닝에 들어 있지 않은 것', '25SS·25FW 러닝 전년 비교(같은 방식으로 세어져 있는지 확인할 자료 없음)'],
    ['', ''],
    ['정의', '‘FW 이월 판매(8/1~)’ = 러닝 ST 시트 누계 − 시점재고(7/31) 시트 누계(날짜 경계 정확). SS 주간판매는 주(월~일) 단위라 8/3 주~ 와 7/27 주~ 두 값의 범위로 본다.'],
    ['연간 러닝', 'SS 러닝 31개(1/1~) + 구품번 추가 %d개(③ FW 구품번 %d + ④ 러닝 원본 %d, SS 스타일맵 12월말 재고+2026 입고 기준 1/1~) + FW 신품번(8/1~). FW 이월은 더하지 않는다.' % (ck['onlyN'], ck['onlyFwN'], ck['onlyExtraN'])],
    ['구품번 추가 합계(SS 기준, ③+④)', '투입 %s장 · 1/1~ 판매 %s장 · 원가 %s원(VAT 미포함) · 12월말 재고+입고−판매−현재고 잔차 최대 %s장' % (format(int(ck['onlyBasis']['qtyIn']), ','), format(int(ck['onlyBasis']['sales']), ','), format(int(ck['onlyBasis']['cost']), ','), int(ck.get('onlyResidMax') or 0))],
    ['구품번 추가 %d개의 1/1~ 판매(주마감 − 12월말 시트) vs 채널표(SS 스타일맵 ‘전주(유통1.1~)’ 1/1~, 날짜 기준)' % ck['onlyN'],
     '%d/%d개 ±%d장 이내 · 합계 %s vs 채널표 %s(차 %s장) · 벗어난 품번: %s · 채널표에 행 없음: %s · 채널표보다 작은 품번: %s · SS 러닝 %d개의 스타일맵 정의 %s vs 채널표 %s(차 %s장)' % (
         ck.get('onlyChN', 0) - len(ck.get('onlyChOut') or []), ck.get('onlyChN', 0), TOL, format(int(ck.get('onlyChBasisSum', 0)), ','), format(int(ck.get('onlyChSum', 0)), ','),
         format(int(ck.get('onlyChBasisSum', 0) - ck.get('onlyChSum', 0)), ','), ', '.join('%s[%s](%s vs %s)' % (x['code'], x['src'], format(int(x['sales']), ','), format(int(x['ch']), ',')) for x in (ck.get('onlyChOut') or [])) or '없음',
         ', '.join(ck.get('onlyChMissing') or []) or '없음', ', '.join(ck.get('onlyChBelow') or []) or '없음', ck['ssN'], format(int(ck.get('ssChBackSum', 0)), ','), format(int(ck.get('ssChSum', 0)), ','), format(int(ck.get('ssChBackSum', 0) - ck.get('ssChSum', 0)), ','))],
    ['  원인(검증 확인)', '12월말 기초재고(런닝) 시트에는 빠져 있고 현재 주마감 누계에는 들어 있는 2025년 일자 판매(8월 말~9월 초 소급 입력으로 보임) — 12월 시트가 그 전에 굳어 있어 ‘주마감 − 12월 시트’가 날짜 기준보다 크다. 투입(= 12월말 재고 + 입고)도 같은 양만큼 크게 잡힌다. 합계에는 스타일맵 값을 그대로 사용(SS 러닝 31개와 같은 정의로 맞추려고).'],
    ['(참고) 이전 방식: 구품번 추가 %d개의 1/1~ 판매 vs FW 주간판매(12/29 주부터 ~ 1/5 주부터 범위, ±%d장)' % (ck['onlyN'], TOL),
     '%d/%d개 범위 안 · 밖: %s — 벗어난 양의 합 %s장' % (ck.get('onlyWeeklyN', 0) - len(ck.get('onlyWeeklyOut') or []), ck.get('onlyWeeklyN', 0),
                              ', '.join('%s[%s](%s장 vs %s~%s)' % (x['code'], x.get('src', ''), format(int(x['sales']), ','), format(int(x['lo']), ','), format(int(x['hi']), ',')) for x in (ck.get('onlyWeeklyOut') or [])) or '없음',
                              format(int(sum(max(x['sales'] - x['hi'], x['lo'] - x['sales'], 0) for x in (ck.get('onlyWeeklyOut') or []))), ','))],
    ['  (참고) 원인', '주간판매 범위 방식은 12월말 기준일이 주 중간이라 범위가 넓어 3% 과대를 못 잡는다 — 위 채널표 대조가 더 정확하다'],
    ['과년 러닝(④) 점검', '%d개 · SS·FW 스타일맵 현재고 최대 차이 %s장 · 복종 분류 불일치 %s · FW 신품번 목록과 겹침 %s · 자료 없어 못 넣은 품번 %s' % (
        ck['onlyExtraN'], int(ck.get('extraStockDiffMax') or 0), ', '.join(ck.get('extraGrpMismatch') or []) or '없음', ', '.join(ck.get('extraInFwNew') or []) or '없음',
        ', '.join(u['code'] for u in (ck.get('unlisted') or [])) or '없음')],
]
for ln in lines:
    ws.append(ln)
ws['A1'].font = Font(bold=True, size=14)
ws.column_dimensions['A'].width = 62
ws.column_dimensions['B'].width = 90
for row in ws.iter_rows(min_row=3):
    row[0].font = Font(bold=row[0].value in ('무엇을 점검했나', '대조 결과(허용 오차 %d장)' % TOL, '정의', '연간 러닝'))
    if isinstance(row[1].value, (int, float)):
        row[1].number_format = '#,##0'
        row[1].alignment = Alignment(horizontal='left')


def ov_rows(lst):
    out = []
    for r in lst:
        out.append([r['code'], r.get('name', ''), r.get('item', ''), r.get('ssIn'), r.get('ssSales'), r.get('ssSalesRec'), (r.get('ssSales') or 0) - (r.get('ssSalesRec') or 0) if r.get('ssSalesRec') is not None else None,
                    r.get('fwCarryStock'), r.get('fwAugSales'), r.get('ssWk0803'), r.get('ssWk0727'),
                    r.get('ssStock'), r.get('fwStock'), (r.get('ssStock') or 0) - (r.get('fwStock') or 0), '○' if r.get('boundaryOk') else '× 확인'])
    return out


hdr = ['품번', '스타일명', '복종', '입고', '1/1~ 판매(스타일맵 정의)', '판매기록 기준(주마감−12월말)', '차이', 'FW 이월 재고(7/31)', 'FW 이월 판매(8/1~)', 'SS 주간판매 8/3주~', 'SS 주간판매 7/27주~', 'SS 현재고', 'FW 현재고', '현재고 차이', '경계 안(±%d)' % TOL]
w13 = [14, 30, 9, 12, 14, 14, 9, 14, 14, 14, 14, 11, 11, 11, 12]
sheet(wb.create_sheet(), '① SS러닝과 겹침', hdr, ov_rows(ov1), w13)
sheet(wb.create_sheet(), '② SS신상품과 겹침', hdr, ov_rows(ov2), w13)
sheet(wb.create_sheet(), '③④ 구품번 추가(SS 기준)', ['품번', '스타일명', '복종', '12월말 재고', '2026 입고', '입고(투입)', '1/1~ 판매', '현재고', '원가(VAT 미포함)', '잔차(재고+입고−판매−현재고)', '참고: FW 이월 재고(7/31)', '참고: FW 이월 판매(8/1~)', 'FW 주간판매 12/29주~', 'FW 주간판매 1/5주~', '주간판매 범위 안(±%d)' % TOL, '출처', '채널표 1/1~ 판매(날짜 기준)', '차이(기록 − 채널표)'],
      [[r['code'], r['name'], r['item'], (r['basis'] or {}).get('base'), (r['basis'] or {}).get('newIn'), (r['basis'] or {}).get('qtyIn'), (r['basis'] or {}).get('sales'), (r['basis'] or {}).get('stockNow'), (r['basis'] or {}).get('cost'), (r['basis'] or {}).get('resid'), r['fwCarryStock'], r['fwAugSales'], r.get('wkHi'), r.get('wkLo'), ('○' if r.get('wkOk') else '× 확인'), r.get('src', '') + ((' (' + '·'.join(r['nasSeasons']) + ' 목록)') if r.get('nasSeasons') else ''), r.get('chSales'), ((r['basis'] or {}).get('sales', 0) - r['chSales']) if r.get('chSales') is not None else None] for r in only], [14, 34, 9, 12, 12, 12, 12, 11, 16, 16, 14, 14, 14, 14, 14, 22, 16, 14])
sheet(wb.create_sheet(), 'SS 러닝 전 품번 대조', ['품번', '겹침(FW 구품번)', '스타일맵 정의 판매', '판매기록 기준 판매', 'SS 주간판매 7/27주~', 'FW 주간판매 7/27주~', '52주 최대 차이', 'SS 현재고', 'FW 현재고', '현재고 차이'],
      [[r['code'], '○' if r['overlap'] else '', r.get('ssSalesBack'), r.get('ssSalesRec'), r['ssWk0727'], r['fwWk0727'], r['weeklyMaxDiff'], r['ssStock'], r['fwStock'], r['ssStock'] - r['fwStock']] for r in allss], [14, 14, 16, 16, 18, 18, 14, 12, 12, 12])
if ck.get('unlisted'):                                                   # SS 스타일맵에 자료가 없어 못 넣은 품번(있을 때만)
    sheet(wb.create_sheet(), '자료 없어 못 넣은 품번', ['품번', '러닝 원본(NAS)', '25년 러닝 목록'], [[u['code'], '○' if u.get('inNas') else '', '○' if u.get('in25') else ''] for u in ck['unlisted']], [14, 14, 14])
wb.save(OUT)
print('저장:', OUT)
